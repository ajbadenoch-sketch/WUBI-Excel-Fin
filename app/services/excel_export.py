import io
from datetime import date, timedelta
from calendar import monthrange
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from sqlalchemy import func

from app.models import (
    User, DailyPlan, Task, WeeklyReview, Transaction, Debt,
    DebtPayment, Account, Budget, HealthMetric, Exercise,
    Habit, HabitLog, TaskPriority, TaskStatus, TransactionType
)

# Styles
HEADER_FILL = PatternFill("solid", fgColor="1a1a2e")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
SECTION_FILL = PatternFill("solid", fgColor="16213e")
SECTION_FONT = Font(name="Calibri", bold=True, color="e94560", size=13)
POSITIVE_FILL = PatternFill("solid", fgColor="d4edda")
NEGATIVE_FILL = PatternFill("solid", fgColor="f8d7da")
NEUTRAL_FILL = PatternFill("solid", fgColor="f0f0f0")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
MXN_FORMAT = '#,##0.00" MXN"'


def style_header_row(ws, row, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER


def export_full_report(user: User, db: Session) -> io.BytesIO:
    wb = Workbook()

    _build_dashboard_sheet(wb, user, db)
    _build_plos_sheet(wb, user, db)
    _build_financial_sheet(wb, user, db)
    _build_health_sheet(wb, user, db)
    _build_debt_sheet(wb, user, db)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _build_dashboard_sheet(wb, user, db):
    ws = wb.active
    ws.title = "Dashboard"
    today = date.today()
    month_start = date(today.year, today.month, 1)
    _, dim = monthrange(today.year, today.month)
    month_end = date(today.year, today.month, dim)

    ws.merge_cells("A1:F1")
    ws["A1"] = f"VidaControl - Reporte Integral - {user.name}"
    ws["A1"].font = Font(name="Calibri", bold=True, size=16, color="1a1a2e")
    ws["A2"] = f"Generado: {today.isoformat()}"

    # Scores section
    row = 4
    ws.cell(row=row, column=1, value="SCORES MAESTROS").font = SECTION_FONT
    row += 1
    headers = ["Metrica", "Score", "Estado"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, 3)

    # Calculate scores
    tasks = db.query(Task).filter(Task.user_id == user.id, Task.plan_date == today).all()
    total_t = len(tasks)
    done_t = len([t for t in tasks if t.status == TaskStatus.DONE])
    exec_s = round((done_t / max(total_t, 1)) * 100)

    income = db.query(func.coalesce(func.sum(Transaction.amount), 0)).filter(
        Transaction.user_id == user.id,
        Transaction.transaction_type == TransactionType.INCOME,
        Transaction.transaction_date >= month_start,
        Transaction.transaction_date <= month_end,
    ).scalar()
    expenses = db.query(func.coalesce(func.sum(Transaction.amount), 0)).filter(
        Transaction.user_id == user.id,
        Transaction.transaction_type == TransactionType.EXPENSE,
        Transaction.transaction_date >= month_start,
        Transaction.transaction_date <= month_end,
    ).scalar()
    net = income - expenses
    fin_s = min(100, max(0, 50 + int(net / max(income, 1) * 50))) if income > 0 else 50

    row += 1
    scores = [
        ("Ejecucion", exec_s),
        ("Finanzas", fin_s),
        ("Salud", 50),
        ("GENERAL", round(exec_s * 0.35 + fin_s * 0.35 + 50 * 0.30)),
    ]
    for label, score in scores:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=score)
        estado = "Excelente" if score >= 80 else "Bien" if score >= 60 else "Mejorar"
        ws.cell(row=row, column=3, value=estado)
        fill = POSITIVE_FILL if score >= 60 else NEGATIVE_FILL
        for c in range(1, 4):
            ws.cell(row=row, column=c).fill = fill
            ws.cell(row=row, column=c).border = THIN_BORDER
        row += 1

    # Financial summary
    row += 1
    ws.cell(row=row, column=1, value="RESUMEN FINANCIERO MES").font = SECTION_FONT
    row += 1
    fin_data = [
        ("Ingresos", income),
        ("Gastos", expenses),
        ("Neto", net),
    ]
    for label, val in fin_data:
        ws.cell(row=row, column=1, value=label)
        cell = ws.cell(row=row, column=2, value=val)
        cell.number_format = MXN_FORMAT
        cell.fill = POSITIVE_FILL if val >= 0 else NEGATIVE_FILL
        row += 1

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 15


def _build_plos_sheet(wb, user, db):
    ws = wb.create_sheet("Sistema Operativo")
    today = date.today()
    ws_start = today - timedelta(days=today.weekday())

    ws.cell(row=1, column=1, value="SISTEMA OPERATIVO PERSONAL").font = SECTION_FONT

    # Weekly reviews
    row = 3
    headers = ["Semana", "Tareas Plan", "Completadas", "%Ejecucion",
               "Mood Prom", "Energia Prom", "Estres Prom", "Score General"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, len(headers))

    reviews = db.query(WeeklyReview).filter(
        WeeklyReview.user_id == user.id,
    ).order_by(WeeklyReview.week_start.desc()).limit(12).all()

    row += 1
    for r in reviews:
        ws.cell(row=row, column=1, value=r.week_start.isoformat())
        ws.cell(row=row, column=2, value=r.tasks_planned or 0)
        ws.cell(row=row, column=3, value=r.tasks_completed or 0)
        ws.cell(row=row, column=4, value=r.execution_score or 0)
        ws.cell(row=row, column=5, value=round(r.avg_mood or 0, 1))
        ws.cell(row=row, column=6, value=round(r.avg_energy or 0, 1))
        ws.cell(row=row, column=7, value=round(r.avg_stress or 0, 1))
        ws.cell(row=row, column=8, value=r.overall_score or 0)
        for c in range(1, len(headers) + 1):
            ws.cell(row=row, column=c).border = THIN_BORDER
        row += 1

    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16


def _build_financial_sheet(wb, user, db):
    ws = wb.create_sheet("Estado Financiero")
    today = date.today()

    ws.cell(row=1, column=1, value="ESTADO DE RESULTADOS PERSONAL").font = SECTION_FONT

    # P&L by month (last 6 months)
    row = 3
    headers = ["Mes", "Ingresos", "Gastos", "Neto", "Margen %"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, len(headers))

    row += 1
    chart_start = row
    for offset in range(5, -1, -1):
        m = today.month - offset
        y = today.year
        while m <= 0:
            m += 12
            y -= 1
        _, dim = monthrange(y, m)
        ms = date(y, m, 1)
        me = date(y, m, dim)

        inc = db.query(func.coalesce(func.sum(Transaction.amount), 0)).filter(
            Transaction.user_id == user.id,
            Transaction.transaction_type == TransactionType.INCOME,
            Transaction.transaction_date >= ms,
            Transaction.transaction_date <= me,
        ).scalar()
        exp = db.query(func.coalesce(func.sum(Transaction.amount), 0)).filter(
            Transaction.user_id == user.id,
            Transaction.transaction_type == TransactionType.EXPENSE,
            Transaction.transaction_date >= ms,
            Transaction.transaction_date <= me,
        ).scalar()
        net = inc - exp
        margin = round((net / inc * 100) if inc > 0 else 0, 1)

        ws.cell(row=row, column=1, value=f"{y}-{m:02d}")
        ws.cell(row=row, column=2, value=inc).number_format = MXN_FORMAT
        ws.cell(row=row, column=3, value=exp).number_format = MXN_FORMAT
        cell_net = ws.cell(row=row, column=4, value=net)
        cell_net.number_format = MXN_FORMAT
        cell_net.fill = POSITIVE_FILL if net >= 0 else NEGATIVE_FILL
        ws.cell(row=row, column=5, value=margin)
        for c in range(1, 6):
            ws.cell(row=row, column=c).border = THIN_BORDER
        row += 1

    # Add chart
    if row > chart_start:
        chart = BarChart()
        chart.title = "Ingresos vs Gastos"
        chart.y_axis.title = "MXN"
        data_ref = Reference(ws, min_col=2, max_col=3, min_row=3, max_row=row - 1)
        cats = Reference(ws, min_col=1, min_row=chart_start, max_row=row - 1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats)
        chart.shape = 4
        ws.add_chart(chart, f"A{row + 1}")

    # Balance sheet
    row += 18
    ws.cell(row=row, column=1, value="BALANCE PERSONAL").font = SECTION_FONT
    row += 1
    headers2 = ["Concepto", "Monto"]
    for i, h in enumerate(headers2, 1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, 2)

    accounts = db.query(Account).filter(
        Account.user_id == user.id, Account.is_active == True
    ).all()
    row += 1
    ws.cell(row=row, column=1, value="ACTIVOS").font = Font(bold=True)
    row += 1
    total_assets = 0
    for a in accounts:
        ws.cell(row=row, column=1, value=f"  {a.name} ({a.account_type})")
        ws.cell(row=row, column=2, value=a.balance).number_format = MXN_FORMAT
        total_assets += a.balance
        row += 1
    ws.cell(row=row, column=1, value="Total Activos").font = Font(bold=True)
    ws.cell(row=row, column=2, value=total_assets).number_format = MXN_FORMAT
    row += 1

    debts = db.query(Debt).filter(
        Debt.user_id == user.id, Debt.is_active == True
    ).all()
    ws.cell(row=row, column=1, value="PASIVOS").font = Font(bold=True)
    row += 1
    total_debt = 0
    for d in debts:
        ws.cell(row=row, column=1, value=f"  {d.name}")
        ws.cell(row=row, column=2, value=d.current_balance).number_format = MXN_FORMAT
        total_debt += d.current_balance
        row += 1
    ws.cell(row=row, column=1, value="Total Pasivos").font = Font(bold=True)
    ws.cell(row=row, column=2, value=total_debt).number_format = MXN_FORMAT
    row += 2
    ws.cell(row=row, column=1, value="PATRIMONIO NETO").font = Font(bold=True, size=12)
    nw_cell = ws.cell(row=row, column=2, value=total_assets - total_debt)
    nw_cell.number_format = MXN_FORMAT
    nw_cell.font = Font(bold=True, size=12)
    nw_cell.fill = POSITIVE_FILL if total_assets >= total_debt else NEGATIVE_FILL

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 14


def _build_health_sheet(wb, user, db):
    ws = wb.create_sheet("Salud")
    today = date.today()

    ws.cell(row=1, column=1, value="SISTEMA DE SALUD").font = SECTION_FONT

    row = 3
    headers = ["Fecha", "Peso (kg)", "Cintura (cm)", "Glucosa Ayunas",
               "Glucosa Post", "PA Sys", "PA Dia", "HbA1c", "Sueno (hrs)",
               "Calidad Sueno", "Estres"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, len(headers))

    metrics = db.query(HealthMetric).filter(
        HealthMetric.user_id == user.id,
    ).order_by(HealthMetric.metric_date.desc()).limit(30).all()

    row += 1
    for m in metrics:
        ws.cell(row=row, column=1, value=m.metric_date.isoformat())
        ws.cell(row=row, column=2, value=m.weight_kg)
        ws.cell(row=row, column=3, value=m.waist_cm)
        glucose_cell = ws.cell(row=row, column=4, value=m.glucose_fasting)
        if m.glucose_fasting and m.glucose_fasting > 130:
            glucose_cell.fill = NEGATIVE_FILL
        elif m.glucose_fasting and m.glucose_fasting <= 100:
            glucose_cell.fill = POSITIVE_FILL
        ws.cell(row=row, column=5, value=m.glucose_post_meal)
        ws.cell(row=row, column=6, value=m.blood_pressure_sys)
        ws.cell(row=row, column=7, value=m.blood_pressure_dia)
        ws.cell(row=row, column=8, value=m.hba1c)
        ws.cell(row=row, column=9, value=m.sleep_hours)
        ws.cell(row=row, column=10, value=m.sleep_quality)
        ws.cell(row=row, column=11, value=m.stress_level)
        for c in range(1, len(headers) + 1):
            ws.cell(row=row, column=c).border = THIN_BORDER
        row += 1

    # Exercise summary
    row += 2
    ws.cell(row=row, column=1, value="EJERCICIO - ULTIMO MES").font = SECTION_FONT
    row += 1
    ex_headers = ["Fecha", "Tipo", "Duracion (min)", "Intensidad"]
    for i, h in enumerate(ex_headers, 1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, len(ex_headers))

    month_ago = today - timedelta(days=30)
    exercises = db.query(Exercise).filter(
        Exercise.user_id == user.id,
        Exercise.exercise_date >= month_ago,
    ).order_by(Exercise.exercise_date.desc()).all()

    row += 1
    for e in exercises:
        ws.cell(row=row, column=1, value=e.exercise_date.isoformat())
        ws.cell(row=row, column=2, value=e.exercise_type.value)
        ws.cell(row=row, column=3, value=e.duration_minutes)
        ws.cell(row=row, column=4, value=e.intensity)
        for c in range(1, 5):
            ws.cell(row=row, column=c).border = THIN_BORDER
        row += 1

    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16


def _build_debt_sheet(wb, user, db):
    ws = wb.create_sheet("Deudas")

    ws.cell(row=1, column=1, value="SISTEMA DE CONTROL DE DEUDA").font = SECTION_FONT

    row = 3
    headers = ["Deuda", "Tipo", "Saldo Original", "Saldo Actual",
               "Tasa %", "Pago Minimo", "% Pagado"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=row, column=i, value=h)
    style_header_row(ws, row, len(headers))

    debts = db.query(Debt).filter(Debt.user_id == user.id).all()

    row += 1
    for d in debts:
        ws.cell(row=row, column=1, value=d.name)
        ws.cell(row=row, column=2, value=d.debt_type.value)
        ws.cell(row=row, column=3, value=d.original_amount).number_format = MXN_FORMAT
        ws.cell(row=row, column=4, value=d.current_balance).number_format = MXN_FORMAT
        ws.cell(row=row, column=5, value=d.interest_rate)
        ws.cell(row=row, column=6, value=d.minimum_payment).number_format = MXN_FORMAT
        paid_pct = round((1 - d.current_balance / max(d.original_amount, 1)) * 100, 1)
        ws.cell(row=row, column=7, value=paid_pct)
        fill = POSITIVE_FILL if paid_pct > 50 else NEUTRAL_FILL
        for c in range(1, len(headers) + 1):
            ws.cell(row=row, column=c).border = THIN_BORDER
            ws.cell(row=row, column=c).fill = fill
        row += 1

    # Totals
    if debts:
        row += 1
        ws.cell(row=row, column=1, value="TOTAL").font = Font(bold=True)
        ws.cell(row=row, column=3, value=sum(d.original_amount for d in debts)).number_format = MXN_FORMAT
        ws.cell(row=row, column=4, value=sum(d.current_balance for d in debts)).number_format = MXN_FORMAT
        ws.cell(row=row, column=6, value=sum(d.minimum_payment for d in debts)).number_format = MXN_FORMAT

    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
